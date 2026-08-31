# -*- coding: utf-8 -*-
"""Rebuild fusion deck on Labillion orange-white template."""
import io
import os
import re
import zipfile
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn, nsmap

SRC = r"c:\Users\Sharon\Desktop\演示文件\化工\镁伽智慧实验室AI解决方案_20260509.pptx"
OUT1 = r"c:\Users\Sharon\Desktop\演示文件\化工\镁伽智慧实验室_调度平台融合方案_20260831.pptx"
OUT2 = r"e:\MEGA\M-技术方案\镁伽智慧实验室_调度平台融合方案_20260831.pptx"
TMP = r"c:\Users\Sharon\Desktop\待整理文件\Chemical-BO\Chemical-BO-main\_labillion_blank.pptx"

ORANGE = RGBColor(0xF0, 0x4E, 0x23)
ORANGE2 = RGBColor(0xF5, 0x6E, 0x23)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x41, 0x41, 0x41)
BLACK = RGBColor(0x00, 0x00, 0x00)
MUTED = RGBColor(0x8D, 0x8D, 0x8D)
CREAM = RGBColor(0xFF, 0xF5, 0xF0)
PEACH = RGBColor(0xFF, 0xE8, 0xDE)
LINE = RGBColor(0xF0, 0xD0, 0xC4)
FONT = "微软雅黑"
NSMAP_REL = {"pr": "http://schemas.openxmlformats.org/package/2006/relationships"}
NSMAP_CT = {"ct": "http://schemas.openxmlformats.org/package/2006/content-types"}
NSMAP_P = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
NSMAP_R = {"r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}


def collect_keep_media(zf):
    keep = set()
    prefixes = (
        "ppt/slideMasters/",
        "ppt/slideLayouts/",
        "ppt/theme/",
        "ppt/notesMasters/",
    )
    for name in zf.namelist():
        if not name.endswith(".rels"):
            continue
        if not any(name.startswith(p.replace("/", "/")) or name.startswith(p) for p in (
            "ppt/slideMasters/_rels/",
            "ppt/slideLayouts/_rels/",
            "ppt/theme/_rels/",
            "ppt/notesMasters/_rels/",
            "ppt/_rels/",
        )):
            continue
        xml = zf.read(name).decode("utf-8", "ignore")
        for m in re.finditer(r'Target="([^"]*media/[^"]+)"', xml):
            tgt = m.group(1)
            # resolve relative
            base = name.rsplit("/", 1)[0].replace("/_rels", "")
            # Target like ../media/image1.png
            parts = (base + "/" + tgt).split("/")
            stack = []
            for p in parts:
                if p == "..":
                    stack.pop()
                elif p and p != ".":
                    stack.append(p)
            keep.add("/".join(stack))
    return keep


def slim_template(src, dst):
    keep_exact_prefix = (
        "[Content_Types].xml",
        "_rels/",
        "docProps/",
        "ppt/_rels/",
        "ppt/slideMasters/",
        "ppt/slideLayouts/",
        "ppt/theme/",
        "ppt/notesMasters/",
        "ppt/presProps.xml",
        "ppt/viewProps.xml",
        "ppt/tableStyles.xml",
        "ppt/commentAuthors.xml",
        "ppt/tags/",
    )
    with zipfile.ZipFile(src, "r") as zin:
        media_keep = collect_keep_media(zin)
        names = set(zin.namelist())
        # presentation.xml + rels handled specially
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
            for name in zin.namelist():
                if name.startswith("ppt/slides/") or name.startswith("ppt/notesSlides/") \
                   or name.startswith("ppt/charts/") or name.startswith("ppt/diagrams/") \
                   or name.startswith("ppt/embeddings/") or name.startswith("ppt/handoutMasters/"):
                    continue
                if name.startswith("ppt/media/") and name not in media_keep:
                    continue
                if name == "ppt/presentation.xml":
                    xml = zin.read(name)
                    root = etree.fromstring(xml)
                    sldIdLst = root.find("{http://schemas.openxmlformats.org/presentationml/2006/main}sldIdLst")
                    if sldIdLst is not None:
                        for child in list(sldIdLst):
                            sldIdLst.remove(child)
                    zout.writestr(name, etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True))
                    continue
                if name == "ppt/_rels/presentation.xml.rels":
                    xml = zin.read(name)
                    root = etree.fromstring(xml)
                    ns = "http://schemas.openxmlformats.org/package/2006/relationships"
                    for rel in list(root):
                        typ = rel.get("Type", "")
                        if typ.endswith("/slide") or typ.endswith("/notesSlide") or typ.endswith("/chart"):
                            root.remove(rel)
                    zout.writestr(name, etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True))
                    continue
                if name == "[Content_Types].xml":
                    xml = zin.read(name)
                    root = etree.fromstring(xml)
                    ns = "http://schemas.openxmlformats.org/package/2006/content-types"
                    for ov in list(root):
                        part = ov.get("PartName", "")
                        if any(x in part for x in ("/ppt/slides/slide", "/ppt/notesSlides/", "/ppt/charts/", "/ppt/embeddings/")):
                            root.remove(ov)
                        if part.startswith("/ppt/media/") and part.lstrip("/") not in media_keep:
                            # Override entries for unused media
                            if ov.tag.endswith("Override"):
                                root.remove(ov)
                    zout.writestr(name, etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True))
                    continue
                if name.startswith(keep_exact_prefix) or name.startswith("ppt/media/"):
                    zout.writestr(name, zin.read(name))
        with open(dst, "wb") as f:
            f.write(buf.getvalue())
    print("slim template bytes", os.path.getsize(dst))


def set_run(run, size=14, bold=False, color=INK, font=FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", font)


def fill_shape_text(shp, text, size=16, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    tf = shp.text_frame
    tf.word_wrap = True
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    # keep first para, drop rest
    p = tf.paragraphs[0]
    p.alignment = align
    p.clear()
    run = p.add_run()
    run.text = text
    set_run(run, size, bold, color)
    # remove extra paragraphs' text
    for extra in tf.paragraphs[1:]:
        extra.clear()


def ph(slide, name):
    for shp in slide.placeholders:
        if shp.name == name:
            return shp
    return None


def hide_placeholders(slide, keep_names=None):
    keep_names = keep_names or set()
    for shp in slide.placeholders:
        if shp.name not in keep_names:
            shp.left = Inches(-15)
            shp.top = Inches(-15)
            shp.width = Inches(1)
            shp.height = Inches(0.3)


def add_rect(slide, l, t, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    sh.shadow.inherit = False
    return sh


def add_round(slide, l, t, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    try:
        sh.adjustments[0] = 0.08
    except Exception:
        pass
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    sh.shadow.inherit = False
    return sh


def add_tb(slide, l, t, w, h, text, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size, bold, color)
    return box


def add_paras(slide, l, t, w, h, lines, size=13, color=INK, spacing=6, bold=False):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(spacing)
        run = p.add_run()
        run.text = line
        set_run(run, size, bold, color)
    return box


def content_slide(prs, title):
    slide = prs.slides.add_slide(prs.slide_layouts[3])  # 内容页
    tshp = ph(slide, "文本占位符 45")
    if tshp:
        fill_shape_text(tshp, title, 26, True, ORANGE, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
    hide_placeholders(slide, keep_names={"文本占位符 45"})
    return slide


def card(slide, l, t, w, h, title, body, orange_head=False):
    if orange_head:
        add_round(slide, l, t, w, h, WHITE, LINE)
        add_rect(slide, l, t, w, Inches(0.48), ORANGE)
        add_tb(slide, l + Inches(0.16), t + Inches(0.06), w - Inches(0.3), Inches(0.38), title, 13, True, WHITE, anchor=MSO_ANCHOR.MIDDLE)
        add_tb(slide, l + Inches(0.16), t + Inches(0.56), w - Inches(0.32), h - Inches(0.68), body, 12, False, INK)
    else:
        add_round(slide, l, t, w, h, WHITE, LINE)
        add_rect(slide, l, t, Inches(0.08), h, ORANGE)
        add_tb(slide, l + Inches(0.22), t + Inches(0.12), w - Inches(0.32), Inches(0.36), title, 14, True, ORANGE)
        add_tb(slide, l + Inches(0.22), t + Inches(0.5), w - Inches(0.36), h - Inches(0.62), body, 12, False, INK)


# ---------- build ----------
slim_template(SRC, TMP)
prs = Presentation(TMP)
LY_COVER_HZ = prs.slide_layouts[0]
LY_TOC = prs.slide_layouts[1]
LY_CHAP = prs.slide_layouts[2]
LY_CONTENT = prs.slide_layouts[3]
LY_END = prs.slide_layouts[4]
LY_COVER_C = prs.slide_layouts[5]

# 1 cover - Hangzhou building template
s = prs.slides.add_slide(LY_COVER_HZ)
add_tb(s, Inches(0.7), Inches(2.15), Inches(11.5), Inches(0.4), "化工智慧实验室  ·  技术方案", 16, False, WHITE)
add_tb(s, Inches(0.7), Inches(2.6), Inches(12), Inches(0.9), "镁伽智慧实验室融合方案", 36, True, WHITE)
add_rect(s, Inches(0.7), Inches(3.6), Inches(2.4), Inches(0.07), ORANGE)
add_tb(s, Inches(0.7), Inches(3.85), Inches(11.5), Inches(0.45), "Labillion 2.0  ×  实验室智能调度平台", 18, False, ORANGE2)
add_tb(s, Inches(0.7), Inches(4.5), Inches(10.5), Inches(0.7), "把「AI 能决策」落到「实验室能跑起来」：统一调度、数据闭环、干湿迭代。", 14, False, WHITE)
add_tb(s, Inches(0.7), Inches(6.55), Inches(10), Inches(0.3), "2026.08.31    MEGAROBO    融合自 Labillion 方案与调度平台 V2.1", 12, False, WHITE)

# 2 TOC
s = prs.slides.add_slide(LY_TOC)
toc_num = {
    "文本占位符 1": "01",
    "文本占位符 5": "02",
    "文本占位符 7": "03",
    "文本占位符 9": "04",
    "文本占位符 2": "05",
    "文本占位符 4": "06",
    "文本占位符 6": "",
}
toc_txt = {
    "文本占位符 10": "目录",
    "文本占位符 3": "融合定位    Labillion 管大脑，调度平台管手脚与现场",
    "文本占位符 18": "镁伽软件特点    大模型 · 智能体 · 柔性调度 · 数采治理",
    "文本占位符 17": "镁伽核心优势    跨实验室统一规划、动态并行、非标 + AGV",
    "文本占位符 16": "同行软件现状    汇像 / 深度原理 / 机数量子 / 深势 / 奔曜",
    "文本占位符 12": "横向对比与选型    各家主战场不同，镁伽卡在可落地调度中枢",
    "文本占位符 11": "落地路径与案例    化工质检、材料干湿闭环、模块化接入",
    "文本占位符 8": "",
}
for shp in s.placeholders:
    if shp.name in toc_num:
        fill_shape_text(shp, toc_num[shp.name], 20, True, ORANGE, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
    elif shp.name in toc_txt:
        sz = 28 if shp.name == "文本占位符 10" else 16
        fill_shape_text(shp, toc_txt[shp.name], sz, shp.name == "文本占位符 10", INK if shp.name != "文本占位符 10" else ORANGE, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)

def chapter(num, title, d1, d2):
    s = prs.slides.add_slide(LY_CHAP)
    mapping = {
        "文本占位符 4": (num, 54, True, ORANGE),
        "文本占位符 10": (title, 28, True, INK),
        "文本占位符 1": (d1, 16, False, MUTED),
        "文本占位符 6": (d2, 16, False, MUTED),
    }
    for shp in s.placeholders:
        if shp.name in mapping:
            t, sz, b, c = mapping[shp.name]
            fill_shape_text(shp, t, sz, b, c, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
    return s

chapter("01", "融合定位", "Labillion 2.0 解决「智」：大模型、智能体、业务流。", "调度平台解决「控」：机台、AGV、工单、样本。两者必须合成一份方案。")

# why fuse
s = content_slide(prs, "为什么要把两份方案合成一份")
cols = [
    ("Labillion 2.0", "自主实验室大模型：私有化、热插拔、信创。\n多智能体：分析 / 视觉 / 真实实验 / 工艺优化。\n智慧管理系统：人机料法环 + 业务流 + BI。\n偏顶层：AI4S、知识、决策、信息化。"),
    ("智能调度平台", "统一调度实验全流程，支持运行中加单。\n多厂商仪器 / 自动化 / AGV 规范接入。\n工单、样本、物料、设备、日志可运营。\n偏现场：把流程真正跑在岛台与跨实验室。"),
    ("融合后的一句话", "上层：AI 设计路线、优化工艺、判定质控。\n中层：动态调度拆到机台 / AGV / 岛区。\n下层：IoT 把孤岛仪器变成可采可控资源。\n交付：可落地的化工智慧实验室操作系统。"),
]
for i, (t, b) in enumerate(cols):
    card(s, Inches(0.45) + Inches(i * 4.2), Inches(1.55), Inches(4.0), Inches(5.15), t, b, orange_head=True)

# pain
s = content_slide(prs, "实验室智能化的真实卡点")
pains = [
    ("顶层设计碎片化", "LIMS / ELN / 原厂调度 / AI 工具各管一段，换工艺就要重做集成，拓展成本高。"),
    ("数据是垃圾不是资产", "仪器、软件、文献互不相通；自动化越大，堰塞湖越大，AI 没有可训练燃料。"),
    ("调度仍是先到先得", "孔板 FIFO 造成空闲等待；管瓶皿非标抓取断点；多中心 AGV 没有一体化排程。"),
    ("AI 找不到锚点", "模型很强但下不到机台；湿实验验证缺失；核心工艺不敢交给外部黑盒。"),
]
for i, (t, d) in enumerate(pains):
    x = Inches(0.45) + Inches((i % 2) * 6.35)
    y = Inches(1.55) + Inches((i // 2) * 2.55)
    card(s, x, y, Inches(6.1), Inches(2.35), f"0{i+1}  {t}", d)

# architecture
s = content_slide(prs, "融合架构：软件层 / 调度层 / 实验室层")
layers = [
    (ORANGE, "软件层  SOFTWARE", "AI 智能体 / 知识库 / 数据大屏    ·    LIMS · ELN · 业务流运营（LibraX）    ·    辅助研发决策、审计与知识沉淀"),
    (ORANGE2, "调度层  SCHEDULING  ← 本方案核心", "动态调度引擎 · 并行实验 · 运行中加单    ·    控制自动化设备 / 分析仪器 / 转运 AGV    ·    7×24 无人值守"),
    (INK, "实验室层  LAB", "LAB1 化学合成：仓储 / 投料 / 反应 / LCMS    ·    LAB2 电池材料：配液 / 理化 / 组装    ·    LAB3 钙钛矿：合成 / 薄膜 / 表征"),
]
for i, (c, t, d) in enumerate(layers):
    y = Inches(1.55) + Inches(i * 1.7)
    add_round(s, Inches(0.45), y, Inches(12.4), Inches(1.52), WHITE, LINE)
    add_rect(s, Inches(0.45), y, Inches(0.14), Inches(1.52), c)
    add_tb(s, Inches(0.85), y + Inches(0.18), Inches(11.7), Inches(0.42), t, 18, True, ORANGE if i < 2 else INK)
    add_tb(s, Inches(0.85), y + Inches(0.7), Inches(11.7), Inches(0.6), d, 13, False, MUTED)

chapter("02", "镁伽软件特点与优势", "一套操作系统，而不是一堆工具。", "以硬件自动化控制为基础，融合 AI 智能体、业务流、调度策略与数据治理。")

# panorama
s = content_slide(prs, "镁伽软件全景")
mods = [
    ("寰宇大模型", "私有化基座、200+ 领域模型热插拔、多模态、统一 AI-native 入口"),
    ("智能体矩阵", "真实实验 / 工艺优化 / 反应设计 / 方法开发 / 谱图 / 视觉 / 分析"),
    ("柔性执行", "Automation 调度：SBS 板式 · 管瓶皿非标 · AGV 跨实验室"),
    ("卓越业务流", "人机料法环、工单、权限、BI 驾驶舱、跨部门项目派发"),
    ("万象感知", "IoT 数采、CDM 统一模型、安全分级、LIMS/ELN/MES 集成"),
    ("仪器智能体", "Auflo / 合成 / 分装 / 前处理等边缘决策单元，形成规模效应"),
]
for i, (t, d) in enumerate(mods):
    x = Inches(0.45) + Inches((i % 3) * 4.2)
    y = Inches(1.5) + Inches((i // 3) * 2.55)
    card(s, x, y, Inches(4.0), Inches(2.35), t, d, orange_head=True)

# scheduling features
s = content_slide(prs, "特点一：为实验室而生的动态调度")
pts = [
    ("动态调度引擎", "多实验并行；运行中实时加单；按目标优化资源，而不是 FIFO 排队。"),
    ("多厂商兼容", "移液机器人、离心机、培养箱、酶标仪等标准化接口，驱动库持续扩展（200+）。"),
    ("三种调度形态", "SBS 孔板高通量；管 / 瓶 / 皿非标耗材；AGV 跨岛区、跨楼层一体化排程。"),
    ("0 代码编排", "流程、方法、参数可视化配置，自动识别逻辑错误，降低自动化门槛。"),
    ("7×24 黑灯运行", "远程监控、告警、异常自愈；稳定超长连接，保障通量而不是演示。"),
    ("跨实验室统一规划", "多岛区一张图：合成、电池、钙钛矿可模块化复用同一套调度能力。"),
]
for i, (t, d) in enumerate(pts):
    x = Inches(0.45) + Inches((i % 3) * 4.2)
    y = Inches(1.5) + Inches((i // 3) * 2.55)
    card(s, x, y, Inches(4.0), Inches(2.35), t, d)

# data
s = content_slide(prs, "特点二：数采到决策的闭环")
card(s, Inches(0.45), Inches(1.5), Inches(6.15), Inches(5.2), "IoT 强覆盖 + 轻量治理",
     "串口 RS-232/485、TCP/Modbus、文件交换、RPA/图像识别、HID 操作模拟。\n\n物理隔离、无网络的孤岛设备也可采集——这是化工老实验室的刚需。\n\n边缘接入 → 任务生命周期 → 解析标准化（CDM）→ 监控告警 → 上层 BI/AI。\n\nAI 解析通讯协议文档，新旧设备分钟级适配。", True)
card(s, Inches(6.8), Inches(1.5), Inches(6.05), Inches(5.2), "科研账本 + 运营账本",
     "科研：配方 / 工艺参数 / 反应条件；LCMS 峰图、光谱、图像、报告归一化存储。\n\n运营：设备状态与告警、工单达成率、工时产出、OEE、设备 ROI。\n\n闭环对象：一切绑定样品 / 批次，可追溯、可复现。\n\n上接 LIMS，下接机台，旁接 AI 模型。", True)

# agents
s = content_slide(prs, "特点三：AI 智能体长在执行链上")
agents = [
    ("真实实验智能体", "自然语言描述需求 → 意图解析 → 知识库配方法 → 脚本下发设备；样本全生命周期监控。"),
    ("工艺优化（EDBO）", "连续 / 离散参数空间组合推进；高通量自动化验证；保结果质量同时保过程效率。"),
    ("化学反应设计", "垂类反应预测 + 反应网络机理 + 干湿数据资产；从产率预判到自适应寻优。"),
    ("分析方法开发", "杂质风险预测 → 初始方法 v0 + DOE → 贝叶斯多目标迭代，少轮次收敛 HPLC 等。"),
    ("谱图 / 分析 / 视觉", "自动积分、杂质库检索、纯度预测；EC50/DMPK；微米级异物视觉闭环。"),
    ("Auflo 工作站", "SOP 编译为 Worklist；持续学习移液精度，减少人工校准与非计划停机。"),
]
for i, (t, d) in enumerate(agents):
    y = Inches(1.42) + Inches(i * 0.85)
    add_round(s, Inches(0.45), y, Inches(12.4), Inches(0.78), WHITE, LINE)
    add_rect(s, Inches(0.45), y, Inches(2.55), Inches(0.78), ORANGE)
    add_tb(s, Inches(0.55), y + Inches(0.15), Inches(2.35), Inches(0.5), t, 12, True, WHITE, anchor=MSO_ANCHOR.MIDDLE)
    add_tb(s, Inches(3.2), y + Inches(0.15), Inches(9.4), Inches(0.5), d, 13, False, INK, anchor=MSO_ANCHOR.MIDDLE)

# nine modules
s = content_slide(prs, "特点四：九大模块，按实验室日常工作设计")
mods = [
    ("01 运行总览", "分析大屏 + 工作台：温湿度、工单、设备、AGV、样本统计"),
    ("02 工单管理", "手动 / 接口建单，状态标签，样本地图跟踪流转"),
    ("03 实验室", "机台表单、流程、方法、结果一体；步骤组件可复用"),
    ("04 AGV 调度", "任务下发、最优车辆、路网地图、孔位路径逐步可视"),
    ("05 物料管理", "库位图形化（空 / 满 / 预定位），物料与载具模型维护"),
    ("06 样本管理", "与检测系统同步，批量导入，检测项 / 方法 / 仪器配置"),
    ("07 设备管理", "就绪 / 忙碌 / 禁用 / 异常；按区域类型筛选"),
    ("08 日志中心", "接口 / 告警 / 操作三类留痕，满足联调与审计"),
    ("09 用户中心", "跨实验室数据权限：角色、菜单、部门、账号"),
]
for i, (t, d) in enumerate(mods):
    x = Inches(0.4) + Inches((i % 3) * 4.2)
    y = Inches(1.45) + Inches((i // 3) * 1.75)
    card(s, x, y, Inches(4.05), Inches(1.6), t, d, orange_head=True)

# vs traditional
s = content_slide(prs, "镁伽软件优势（相对传统栈）")
headers = ["对照对象", "他们强在哪", "卡点", "镁伽怎么补"]
rows = [
    ["传统 LIMS/ELN", "人机料法环、记录成熟", "执行弱、结构化弱、个性化贵", "记录、执行、调度打通，样本进结果出"],
    ["原厂调度软件", "单岛台通量高", "跨品牌复杂工作流难集成", "多厂商 + SBS/非标/AGV 统一编排"],
    ["MES/LES", "能对到设备", "操作重，常需再买 LIMS/ELN", "实验室科研场景原生，不是工厂套用"],
    ["云数据平台", "存得下", "设备端采集弱、缺治理", "边缘数采 + CDM + 科研/运营双账本"],
    ["CADD/AIDD 工具", "设计快", "缺湿实验验证，算力成本高", "设计结果下发调度，湿实验回流训练"],
]
add_rect(s, Inches(0.4), Inches(1.45), Inches(12.5), Inches(0.46), ORANGE)
cw = 3.12
for i, h in enumerate(headers):
    add_tb(s, Inches(0.5) + Inches(i * cw), Inches(1.5), Inches(3.0), Inches(0.36), h, 13, True, WHITE, anchor=MSO_ANCHOR.MIDDLE)
for r, row in enumerate(rows):
    y = Inches(1.95) + Inches(r * 0.9)
    bg = WHITE if r % 2 == 0 else CREAM
    add_rect(s, Inches(0.4), y, Inches(12.5), Inches(0.88), bg)
    for c, txt in enumerate(row):
        add_tb(s, Inches(0.5) + Inches(c * cw), y + Inches(0.12), Inches(3.0), Inches(0.64), txt, 12, c == 0, ORANGE if c == 0 else INK)

chapter("03", "同行软件现状", "公开信息整理（2025–2026）。", "同场竞技的是「实验室智能化」，真正重叠的层并不一样。")

# landscape
s = content_slide(prs, "同行格局：五家不是同一类产品")
add_tb(s, Inches(0.5), Inches(1.4), Inches(12.3), Inches(0.7),
       "市场拆成三层：① 设备中控 / 调度操作系统    ② 机器人 + 工作站全栈    ③ AI for Science 大脑。镁伽融合方案覆盖 ① 和部分 ③，并以调度把 ② 的硬件纳入。",
       13, False, MUTED)
firms = [
    ("汇像", "中控 OS + 机器人全栈", "检验检测 / 生科 / 食化环"),
    ("深度原理", "生成式 AI + 第一性原理", "新材料 / 新能源 / 日化"),
    ("机数量子", "机器化学家 + 材料库", "催化剂 / 电池 / 半导体"),
    ("深势科技", "开源实验 OS + 计算闭环", "干湿闭环科研基础设施"),
    ("奔曜科技", "生科自动化排程软件", "生命科学高通量工作站"),
]
for i, (n, p, sc) in enumerate(firms):
    x = Inches(0.4) + Inches(i * 2.56)
    add_round(s, x, Inches(2.25), Inches(2.42), Inches(4.35), WHITE, LINE)
    add_rect(s, x, Inches(2.25), Inches(2.42), Inches(1.15), ORANGE)
    add_tb(s, x + Inches(0.08), Inches(2.45), Inches(2.26), Inches(0.8), n, 16, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    add_tb(s, x + Inches(0.12), Inches(3.6), Inches(2.18), Inches(1.3), p, 13, True, ORANGE, PP_ALIGN.CENTER)
    add_tb(s, x + Inches(0.12), Inches(5.05), Inches(2.18), Inches(1.2), sc, 12, False, MUTED, PP_ALIGN.CENTER)

def competitor(title, left_lines, r1t, r1b, r2t, r2b):
    s = content_slide(prs, title)
    add_round(s, Inches(0.4), Inches(1.45), Inches(8.25), Inches(5.25), WHITE, LINE)
    add_paras(s, Inches(0.65), Inches(1.65), Inches(7.85), Inches(4.9), left_lines, 13, INK, 8)
    card(s, Inches(8.85), Inches(1.45), Inches(4.05), Inches(2.5), r1t, r1b, True)
    card(s, Inches(8.85), Inches(4.15), Inches(4.05), Inches(2.55), r2t, r2b, True)

competitor(
    "汇像科技（X-Imaging）",
    [
        "产品：iMagicOS 智慧实验室数字化操作系统 / 中控引擎",
        "公开能力：分布式架构；驱动、调度、监控、追踪、管理；宣称 1000+ 仪器驱动。",
        "硬件侧：HelenX 协作 / 移动机器人、工作站；视觉 AI。",
        "上层：科学大模型；任务调度 + 数据采集 + 合规记录（21 CFR Part 11）。",
        "生态：参编黑灯实验室评价规范；公开称落地 400+ 标杆场景。",
        "主战场：生命科学、临床诊断、样品检测、食化环。",
    ],
    "对镁伽的启示",
    "驱动库规模与检测合规是他们的牌面。比拼设备清单意义有限，要打化工跨实验室动态调度 + 非标耗材 + 工艺 / 质控智能体。",
    "公开信息中的空隙",
    "化工多岛协同、运行中加单、管瓶皿非标、与存量 LIMS 深度闭环的产品化细节较少。中控强，不等于目标优化调度强。",
)

competitor(
    "深度原理（Deep Principle）",
    [
        "定位：MIT 背景团队；生成式 AI + 第一性原理，服务化学 / 材料研发。",
        "软件：ReactiveAI——ReactGen / Reactify / ReactControl / ReactBO / ReactNet / ReactHTE。",
        "智能体：Agent Mira，自然语言驱动调研—设计—计算—验证。",
        "设施：AI Materials Factory，宣传 L4 高通量自主实验室，ECML 闭环。",
        "落地方向：新材料、新能源、精细化工、日化配方。",
        "卖点是「发现引擎」，不是多实验室运营系统。",
    ],
    "重叠与区隔",
    "工艺优化 / 反应网络 / 贝叶斯实验设计，概念层与 Labillion 智能体接近。他们卖发现，镁伽卖可调度执行。",
    "客户会怎么选",
    "已有大量异构仪器、要接 LIMS、要跨楼层 AGV、要质检工单——这不是 Mira 的主场。",
)

competitor(
    "机数量子（合肥）",
    [
        "背景：孵化自中科大精准智能化学实验室；量子化学计算 + 大数据 + AI 预测。",
        "平台：机器化学家；Chem-GPT 化学问答与实验建议。",
        "数据：机数大材库 dcaiku，公开宣传千万级化合物与反应路径资产。",
        "硬件：多台科研机器人 + 智能化学工作站。",
        "产业方向：煤化工催化剂、碳中和、聚丙烯等公开案例。",
        "近期专利：多智能体科研平台（技能引擎、子智能体并行、浏览器自动化）。",
    ],
    "强项",
    "垂直化学闭环与材料知识图谱深。适合「按他们的工作站重建一条机器化学家产线」。",
    "和调度平台的差",
    "通用工单 / 权限 / 多实验室统一规划、存量品牌仪器混连、化工质检 LIMS 闭环，不是其公开产品主线。",
)

s = content_slide(prs, "另外两家必须知道：深势、奔曜")
card(s, Inches(0.4), Inches(1.5), Inches(6.2), Inches(5.2), "深势科技  ·  Uni-Lab-OS / 玻尔·跃迁",
     "Uni-Lab-OS：开源、AI-native 实验室 OS；ROS 2/DDS；设备抽象 A/R/A&R。\n\n玻尔·跃迁实验室：商业产品，宣称 1800+ 仪器型号即插即用；自然语言管试剂 / 设备 / 模板实验。\n\n叙事：计算—实验—数据—计算，补 ELN/LIMS 不管执行的短板。\n\n对镁伽：开源连接层会拉低「能连设备」的门槛；差异要打在化工现场调度、工单运营、AGV 与非标、质检合规交付。", True)
card(s, Inches(6.8), Inches(1.5), Inches(6.1), Inches(5.2), "奔曜科技  ·  Bioyond Studio / LabMind",
     "定位：生命科学实验室自动化软件；拖拽流程、仪器库、一键启动。\n\n排程：前瞻排程 + 动态调度；瓶颈设备识别；秒级工艺时间、孔位级控制。\n\n配套：数字孪生预演；对接 LIMS；21 CFR Part 11。\n\n对镁伽：排程算法在生科孔板场景很成熟。化工侧管瓶皿、跨实验室、工艺 / 质检智能体仍是镁伽主场。", True)

chapter("04", "横向对比与落地", "先认清主战场，再谈谁强谁弱。", "验收不看 PPT 层数，看任务能否并行加单、样本能否一盘到底。")

# comparison table
s = content_slide(prs, "横向对比（公开能力口径，非实测排名）")
headers = ["厂商", "核心定位", "调度 / 中控", "AI 大脑", "更适合的客户"]
data = [
    ["镁伽 MEGAROBO", "软硬一体实验室 OS", "动态并行、加单\nSBS/非标/AGV", "智能体长在执行链", "化工/材料存量实验室"],
    ["汇像", "无人化 OS + 机器人", "iMagicOS 全域中控", "视觉 + 科学大模型", "检测/生科/食化环"],
    ["深度原理", "AI for Science 发现", "自有 L4 工厂调度", "Mira + ReactiveAI", "配方/材料发现项目"],
    ["机数量子", "机器化学家 + 材料库", "自有工作站集群", "Chem-GPT + 知识图谱", "重建机器化学家产线"],
    ["深势", "开源实验 OS + 计算", "Uni-Lab 连接层", "玻尔计算/科学智能体", "计算强、要干湿接口"],
    ["奔曜", "生科自动化排程", "前瞻+动态、孔位级", "视觉/机器人偏执行", "生科高通量 GxP 产线"],
]
colw = [2.15, 2.45, 2.55, 2.45, 2.55]
left0 = 0.4
add_rect(s, Inches(left0), Inches(1.42), Inches(sum(colw) + 0.1), Inches(0.42), ORANGE)
acc = left0
for i, h in enumerate(headers):
    add_tb(s, Inches(acc + 0.04), Inches(1.46), Inches(colw[i] - 0.06), Inches(0.34), h, 11, True, WHITE, anchor=MSO_ANCHOR.MIDDLE)
    acc += colw[i]
for r, row in enumerate(data):
    y = 1.88 + r * 0.78
    bg = PEACH if r == 0 else (WHITE if r % 2 else CREAM)
    add_rect(s, Inches(left0), Inches(y), Inches(sum(colw) + 0.1), Inches(0.76), bg)
    acc = left0
    for c, txt in enumerate(row):
        add_tb(s, Inches(acc + 0.04), Inches(y + 0.06), Inches(colw[c] - 0.06), Inches(0.64), txt, 10, c == 0 or r == 0, ORANGE if r == 0 else INK)
        acc += colw[c]

# differentiation
s = content_slide(prs, "镁伽差异化：一句话怎么对外讲")
msgs = [
    ("对采购 / 信息化", "不是再买一套 LIMS，而是在 LIMS 之下补上缺失的执行中枢：工单、机台、AGV、样本、日志、权限一套账。"),
    ("对实验室主任", "运行中可以加单，多实验并行，设备不再按 FIFO 空转；管瓶皿和孔板都能编进同一流程。"),
    ("对科学家", "自然语言和优化算法给出下一组条件，调度平台负责今晚就在岛台上跑，结果自动回流。"),
    ("对集团 / 多实验室", "合成、电池、钙钛矿能力模块化复用；新实验室接入是配置问题，不是再做一个项目。"),
]
for i, (t, d) in enumerate(msgs):
    y = Inches(1.45) + Inches(i * 1.28)
    add_round(s, Inches(0.45), y, Inches(12.4), Inches(1.15), WHITE, LINE)
    add_rect(s, Inches(0.45), y, Inches(2.55), Inches(1.15), ORANGE)
    add_tb(s, Inches(0.58), y + Inches(0.3), Inches(2.3), Inches(0.55), t, 13, True, WHITE, anchor=MSO_ANCHOR.MIDDLE)
    add_tb(s, Inches(3.2), y + Inches(0.22), Inches(9.4), Inches(0.72), d, 14, False, INK, anchor=MSO_ANCHOR.MIDDLE)

# cases
s = content_slide(prs, "已验证场景：质检运营 + 材料干湿闭环")
card(s, Inches(0.4), Inches(1.5), Inches(6.2), Inches(5.2), "化工质检实验室",
     "痛点：有 LIMS 仍人工录入；20+ 样本盘样式；质控规则繁杂。\n\n做法：任务下发 / 结果回传闭环；样本地图全生命周期；质控样自动判定与长期趋势。\n\n调度角色：多盘类型流转、分析仪器对接、进度数字孪生。\n\n价值：样本进、结果出；检测进度可视；人员设备利用率可算。", True)
card(s, Inches(6.8), Inches(1.5), Inches(6.1), Inches(5.2), "材料研发智御实验室",
     "痛点：文献散、手工通量低、工艺逼近经验天花板。\n\n做法：知识库 + 算法选路线；合成-精馏-配液-检测自动化；数据全链路整合。\n\n调度角色：干实验建议变成湿实验工单，过程参数连续采集。\n\n公开结果口径：首轮迭代相关指标提升约 10%（原方案案例，对外请脱敏）。", True)

# path
s = content_slide(prs, "建议落地路径")
steps = [
    ("阶段 A  连得上", "先把关键机台、分析仪器、孤岛数采接进调度层；设备状态与日志可见。"),
    ("阶段 B  跑得动", "工单 + 流程 / 方法 + 样本；单实验室闭环；对接现有 LIMS。已有单岛调度则走接口同步。"),
    ("阶段 C  调得优", "动态并行与加单、AGV 跨实验室、物料库位；OEE / ROI 看板。"),
    ("阶段 D  变得智", "工艺优化 / 方法开发 / 谱图智能体接到调度参数通道；干湿数据回流私有模型。"),
]
for i, (t, d) in enumerate(steps):
    x = Inches(0.4) + Inches(i * 3.2)
    card(s, x, Inches(1.5), Inches(3.05), Inches(4.0), t, d, True)
add_tb(s, Inches(0.5), Inches(5.7), Inches(12.3), Inches(1.0),
       "两种接入：① 现场只有设备 → 直接对接并管理流程；② 单实验室已串起来 → 用软件接口把数据同步进整体平台，保护已有投资。",
       13, False, MUTED)

# talk
s = content_slide(prs, "对外三分钟话术")
add_round(s, Inches(0.45), Inches(1.5), Inches(12.4), Inches(5.2), WHITE, LINE)
add_paras(s, Inches(0.75), Inches(1.7), Inches(11.9), Inches(4.8), [
    "1. 智慧实验室不是再买一个大模型。缺的是把「今晚这批实验」拆到机台、AGV、样本盘上，并且允许中途加单。",
    "2. 汇像把检测无人化和中控 OS 做得更完整；深度原理和机数量子把材料发现的「大脑」做得更垂直；深势把设备连接标准和计算闭环做得更开放；奔曜把生科孔板排程做得更细。",
    "3. 镁伽做的是化工 / 材料存量实验室最难的那一层：多厂商、多岛区、孔板 + 管瓶皿、跨实验室转运、LIMS 已在、AI 还要回流。",
    "4. Labillion 提供智能体与业务运营，调度平台提供可交付的九大模块。两者叠在一起，才是可签约、可验收的智慧实验室。",
    "5. 验收盯三件事：任务能否并行加单、样本能否一盘到底、AI 下一组参数能否在当班被执行并写回。",
], 14, INK, 12)

# end
s = prs.slides.add_slide(LY_END)

prs.save(OUT1)
try:
    prs.save(OUT2)
except Exception as e:
    print("second save fail", e)
print("slides", len(prs.slides))
print("saved", OUT1, os.path.getsize(OUT1))
if os.path.exists(TMP):
    os.remove(TMP)
